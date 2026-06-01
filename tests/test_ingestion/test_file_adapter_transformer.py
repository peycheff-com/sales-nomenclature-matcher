from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pytest

from matcher.ingestion import file_adapter
from matcher.ingestion.base import RawCatalogItem
from matcher.ingestion.transformer import transform_item


class TestFileAdapter:
    def test_detect_delimiter_prefers_tabs_then_semicolon_then_comma(self):
        assert file_adapter._detect_delimiter("a\tb\n1\t2") == "\t"
        assert file_adapter._detect_delimiter("a;b\n1;2") == ";"
        assert file_adapter._detect_delimiter("a,b\n1,2") == ","

    def test_map_columns_normalizes_known_headers(self):
        mapping = file_adapter._map_columns([" Код ", "Наименование", "unknown", "ЕД."])

        assert mapping == {0: "code", 1: "name", 3: "unit"}

    def test_row_to_item_converts_numeric_fields_and_skips_bad_values(self):
        col_map = {
            0: "name",
            1: "weight_value",
            2: "volume_value",
            3: "size_value",
            4: "article",
        }

        item = file_adapter._row_to_item(["Bearing", "1,5", "bad", "2.75", " 6205 "], col_map)

        assert item is not None
        assert item.name == "Bearing"
        assert item.weight_value == 1.5
        assert item.volume_value is None
        assert item.size_value == 2.75
        assert item.article == "6205"
        assert item.product_id is not None
        assert item.product_id.startswith("prd_")

    def test_row_to_item_returns_none_without_name(self):
        assert file_adapter._row_to_item(["", "6205"], {0: "name", 1: "article"}) is None

    def test_parse_csv_handles_semicolon_file_and_generated_product_id(self, tmp_path: Path):
        path = tmp_path / "catalog.csv"
        path.write_text(
            "Код;Наименование;Артикул;Бренд;Вес\n"
            "C1;Подшипник SKF;6205;SKF;1,25\n"
            "C2;;empty-name;SKF;2\n",
            encoding="utf-8",
        )

        items = file_adapter.parse_csv(path)

        assert len(items) == 1
        assert items[0].code == "C1"
        assert items[0].name == "Подшипник SKF"
        assert items[0].article == "6205"
        assert items[0].brand == "SKF"
        assert items[0].weight_value == 1.25
        assert items[0].product_id is not None

    def test_parse_csv_returns_empty_for_empty_file(self, tmp_path: Path):
        path = tmp_path / "empty.csv"
        path.write_bytes(b"")

        assert file_adapter.parse_csv(path) == []

    def test_parse_csv_rejects_unknown_headers(self, tmp_path: Path):
        path = tmp_path / "bad.csv"
        path.write_text("foo,bar\n1,2\n", encoding="utf-8")

        with pytest.raises(ValueError, match="No recognized columns"):
            file_adapter.parse_csv(path)

    def test_parse_xlsx_handles_supported_headers_and_empty_workbook(self, tmp_path: Path):
        path = tmp_path / "catalog.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["product_id", "name", "full_name", "brand", "volume_value"])
        ws.append(["p1", "Pump", "Pump full", "Grundfos", "3,5"])
        ws.append(["p2", None, "No name", "Brand", "1"])
        wb.save(path)
        wb.close()

        items = file_adapter.parse_xlsx(path)

        assert len(items) == 1
        assert items[0].product_id == "p1"
        assert items[0].full_name == "Pump full"
        assert items[0].volume_value == 3.5

        empty_path = tmp_path / "empty.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.delete_rows(1, ws.max_row)
        wb.save(empty_path)
        wb.close()

        assert file_adapter.parse_xlsx(empty_path) == []

    def test_parse_xlsx_returns_empty_when_workbook_has_no_active_sheet(self, tmp_path: Path):
        workbook = SimpleNamespace(active=None)

        with patch("matcher.ingestion.file_adapter.openpyxl.load_workbook", return_value=workbook):
            assert file_adapter.parse_xlsx(tmp_path / "no-active.xlsx") == []

    def test_parse_xlsx_rejects_unknown_headers(self, tmp_path: Path):
        path = tmp_path / "bad.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["foo", "bar"])
        ws.append(["1", "2"])
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="No recognized columns"):
            file_adapter.parse_xlsx(path)

    def test_parse_file_dispatches_by_extension(self, tmp_path: Path):
        csv_path = tmp_path / "catalog.tsv"
        csv_path.write_text("name\tarticle\nItem\tA1\n", encoding="utf-8")

        xlsx_path = tmp_path / "catalog.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name"])
        ws.append(["Item"])
        wb.save(xlsx_path)
        wb.close()

        assert file_adapter.parse_file(csv_path)[0].name == "Item"
        assert file_adapter.parse_file(xlsx_path)[0].name == "Item"

        with pytest.raises(ValueError, match="Unsupported file format"):
            file_adapter.parse_file(tmp_path / "catalog.json")


class TestTransformer:
    def test_transform_item_normalizes_optional_fields_and_builds_search_document(self):
        raw = RawCatalogItem(
            product_id="prd_temp",
            code="C1",
            article="6205",
            name="Подшипник SKF",
            full_name="Подшипник SKF 6205 полный",
            brand="SKF",
            manufacturer="SKF Group",
            manufacturer_code="M6205",
            category_id="cat_1",
            category_path="Bearings",
            unit="pcs",
            packaging="box",
            size_value=20.0,
            size_unit="mm",
            weight_value=1.2,
            weight_unit="kg",
            volume_value=0.3,
            volume_unit="l",
            attributes={"color": "steel"},
        )

        def fake_pipeline(text: str):
            return SimpleNamespace(
                text=f"norm:{text}", brand="brand:norm" if text == "SKF" else None
            )

        with patch("matcher.ingestion.transformer.run_pipeline", side_effect=fake_pipeline):
            result = transform_item(raw)

        assert result["product_id"].startswith("prd_")
        assert result["product_id"] != "prd_temp"
        assert result["normalized_name"] == "norm:Подшипник SKF"
        assert result["normalized_full_name"] == "norm:Подшипник SKF 6205 полный"
        assert result["normalized_brand"] == "brand:norm"
        assert result["attributes_json"] == {"color": "steel"}
        assert result["is_active"] is True
        assert "SKF Group" in result["search_document"]
        assert len(result["source_hash"]) == 16

    def test_transform_item_keeps_onec_product_id_and_handles_minimal_fields(self):
        raw = RawCatalogItem(
            product_id="onec-native-id",
            onec_ref="onec-ref",
            name="Item",
        )

        with patch(
            "matcher.ingestion.transformer.run_pipeline",
            return_value=SimpleNamespace(text="norm:item", brand=None),
        ):
            result = transform_item(raw)

        assert result["product_id"] == "onec-native-id"
        assert result["onec_ref"] == "onec-ref"
        assert result["normalized_full_name"] is None
        assert result["normalized_brand"] is None
        assert result["search_document"] == "Item"
        assert result["attributes_json"] == {}

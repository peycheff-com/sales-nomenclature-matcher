from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import openpyxl
import pandas as pd
import pytest

from matcher.ingestion import parser


def _xlsx_bytes(rows: list[list], *, second_sheet_rows: list[list] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Small"
    for row in rows:
        ws.append(row)
    if second_sheet_rows is not None:
        ws2 = wb.create_sheet("Big")
        for row in second_sheet_rows:
            ws2.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


class TestParserReadersAndHeuristics:
    def test_read_csv_raw_handles_bom_semicolon_and_fallback_delimiter(self):
        df = parser._read_csv_raw("\ufeffname;unit\nPump;pcs\n".encode())

        assert df.iloc[0, 0] == "name"
        assert df.iloc[1, 1] == "pcs"

        with patch("matcher.ingestion.parser.csv.Sniffer") as Sniffer:
            Sniffer.return_value.sniff.side_effect = parser.csv.Error
            df = parser._read_csv_raw(b"name\tunit\nPump\tpcs\n")

        assert list(df.iloc[1]) == ["Pump", "pcs"]

        with patch("matcher.ingestion.parser.csv.Sniffer") as Sniffer:
            Sniffer.return_value.sniff.side_effect = parser.csv.Error
            df = parser._read_csv_raw(b"name;unit\nPump;pcs\n")
        assert list(df.iloc[1]) == ["Pump", "pcs"]

        with patch("matcher.ingestion.parser.csv.Sniffer") as Sniffer:
            Sniffer.return_value.sniff.side_effect = parser.csv.Error
            df = parser._read_csv_raw(b"name,unit\nPump,pcs\n")
        assert list(df.iloc[1]) == ["Pump", "pcs"]

    def test_read_csv_raw_falls_back_when_decode_or_read_fails(self):
        with (
            patch(
                "matcher.ingestion.parser.chardet.detect",
                return_value={"encoding": "bad-codec"},
            ),
            patch("matcher.ingestion.parser.pd.read_csv") as read_csv,
        ):
            read_csv.side_effect = [RuntimeError("bad parse"), pd.DataFrame([["a", "b"]])]
            df = parser._read_csv_raw(b"a,b\n")

        assert df.iloc[0, 0] == "a"

    def test_read_excel_raw_handles_merged_cells_and_pandas_fallback(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.merge_cells("A1:B1")
        ws["A1"] = "Title"
        ws.append(["name", "unit"])
        ws.append(["Pump", "pcs"])
        buffer = BytesIO()
        wb.save(buffer)
        wb.close()

        df = parser._read_excel_raw(buffer.getvalue())

        assert df.iloc[0, 0] == "Title"
        assert df.iloc[0, 1] == "Title"

        with patch("matcher.ingestion.parser.pd.read_excel", return_value=pd.DataFrame([["x"]])):
            df = parser._read_excel_raw(b"not-an-xlsx")
        assert df.iloc[0, 0] == "x"

    def test_get_best_sheet_selects_most_populated_sheet_and_fallback(self):
        contents = _xlsx_bytes(
            [["name"], ["Small item"]],
            second_sheet_rows=[["name", "unit"], ["Big item", "pcs"], ["Other", "pcs"]],
        )

        df, name = parser._get_best_sheet(contents)

        assert name == "Big"
        assert "Big item" in df.to_string()

        with patch("matcher.ingestion.parser._read_excel_raw", return_value=pd.DataFrame([["x"]])):
            df, name = parser._get_best_sheet(contents)
        assert name in {"Small", "Big"}

        with (
            patch("openpyxl.load_workbook", side_effect=RuntimeError("bad workbook")),
            patch(
                "matcher.ingestion.parser._read_excel_raw",
                return_value=pd.DataFrame([["fallback"]]),
            ),
        ):
            df, name = parser._get_best_sheet(contents)
        assert name == "Sheet1"
        assert df.iloc[0, 0] == "fallback"

        with patch(
            "matcher.ingestion.parser._read_excel_raw",
            side_effect=[RuntimeError("bad sheet"), pd.DataFrame([["best"], ["row"]])],
        ):
            df, name = parser._get_best_sheet(contents)
        assert name == "Big"
        assert df.iloc[0, 0] == "best"

    def test_gap_group_header_and_column_detection(self):
        df = pd.DataFrame(
            [
                ["Номенклатура клиента = ООО Ромашка", None, None, "Каталог", None],
                ["Номенклатура", "Ед. изм", None, "Наименование", "Цена"],
                ["Насос длинное описание", "шт", None, "Эталонный насос", "1200"],
                ["Кабель длинное описание", "м", None, "Эталонный кабель", "900"],
            ]
        )

        assert parser._find_empty_col_gaps(df) == [2]
        assert parser._find_empty_row_gaps(pd.DataFrame([[None, None], ["x", None]])) == [0]
        assert parser._split_into_column_groups(df) == [(0, 1), (3, 4)]
        assert parser._detect_header_row(df, 0, 1) == 1
        assert parser._classify_column(df, 0, 1) == "name"
        assert parser._classify_column(df, 1, 1) == "unit"
        assert parser._classify_column(df, 4, 1) == "price"
        assert parser._classify_column(pd.DataFrame([["Артикул"], ["ABC-1"]]), 0, 0) == "article"
        assert parser._classify_table_role(df, 0, 1, 1) == "supplier_input"
        assert parser._classify_table_role(df, 3, 4, 1) == "catalog_reference"
        assert parser._extract_supplier_name(df, 0, 1, 1) == "ООО Ромашка"
        supplier_df = pd.DataFrame([[None, "Поставщик: Acme"], ["Name", "Unit"]])
        assert (
            parser._extract_supplier_name(supplier_df, 0, 1, 1) == "Acme"
        )

    def test_column_classification_uses_data_content_fallbacks(self):
        df = pd.DataFrame(
            [
                ["h1", "h2", "h3", "h4"],
                ["1", "1000", "Very long product description", "ABC123"],
                ["2", "2500", "Another long product description", "DEF456"],
            ]
        )

        assert parser._classify_column(df, 0, 0) == "number"
        assert parser._classify_column(df, 1, 0) == "price"
        assert parser._classify_column(df, 2, 0) == "name"

        short_df = pd.DataFrame([["h1", "h2", "h3"], ["", "kg", "A"], [None, "m", "ABCDEFGH"]])
        assert parser._classify_column(short_df, 0, 0) == "other"
        assert parser._classify_column(short_df, 1, 0) == "unit"
        assert parser._classify_column(short_df, 2, 0) == "article"

        medium_df = pd.DataFrame([["h"], ["Medium"], ["Longer"]])
        assert parser._classify_column(medium_df, 0, 0) == "other"

    def test_detect_header_row_skips_blank_and_merged_title_rows(self):
        df = pd.DataFrame(
            [
                [None, None],
                ["Merged title", "Merged title"],
                ["name", "unit"],
                ["Long product description", "pcs"],
            ]
        )

        assert parser._detect_header_row(df, 0, 1) == 2

    def test_extract_region_items_deduplicates_and_keeps_unit_price(self):
        df = pd.DataFrame(
            [
                ["Name", "Unit", "Price"],
                ["Pump", "pcs", "100"],
                ["Pump", "pcs", "100"],
                ["Итого", "", "200"],
                ["Cable", "m", None],
                ["", "", ""],
            ]
        )

        items = parser._extract_region_items(df, 0, 2, 0, 0, 1, 2)

        assert [item["raw_text"] for item in items] == ["Pump", "Cable"]
        assert items[0]["unit"] == "pcs"
        assert items[0]["price"] == "100"

        duplicate_header_items = parser._extract_region_items(
            pd.DataFrame([["Name", "Name"], ["Pump", "Backup"]]), 0, 1, 0, 0
        )
        assert duplicate_header_items[0]["original_row"] == {"Name": "Pump", "Name_1": "Backup"}

        assert parser._extract_region_items(df, 0, 2, 0, 4) == []
        assert parser._extract_region_items(pd.DataFrame(), 0, 0, 0, 0) == []

    def test_col_letter_to_index(self):
        assert parser._col_letter_to_index("A") == 0
        assert parser._col_letter_to_index("Z") == 25
        assert parser._col_letter_to_index("AA") == 26


class TestParserAnalysis:
    @pytest.mark.asyncio
    async def test_analyze_structure_via_llm_success_invalid_and_failure(self):
        rows = [["Name", "Unit"], ["Pump", "pcs"]]
        response = SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(
                        content='{"tables":[{"role":"supplier_input"}],"supplier_name":"Acme"}'
                    )
                )
            ]
        )
        client = SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(return_value=response)))
        )

        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", None)),
        ):
            result = await parser._analyze_structure_via_llm(rows)

        assert result["supplier_name"] == "Acme"

        bad_response = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content='{"tables":[]}'))]
        )
        client.chat.completions.create = AsyncMock(return_value=bad_response)
        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", {})),
        ):
            assert await parser._analyze_structure_via_llm(rows) is None

        client.chat.completions.create = AsyncMock(side_effect=RuntimeError("llm down"))
        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", {})),
        ):
            assert await parser._analyze_structure_via_llm(rows) is None

        wide_rows = [["x"] * 27]
        client.chat.completions.create = AsyncMock(return_value=bad_response)
        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", {})),
        ):
            assert await parser._analyze_structure_via_llm(wide_rows) is None
        prompt = client.chat.completions.create.await_args.kwargs["messages"][0]["content"]
        assert "AA" in prompt

        with patch("matcher.ingestion.parser.llm_available", return_value=False):
            assert await parser._analyze_structure_via_llm(rows) is None

    @pytest.mark.asyncio
    async def test_identify_column_via_llm_success_and_fallbacks(self):
        df = pd.DataFrame({"Code": ["1"], "Name": ["Pump"]})
        response = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content='{"target_column":"Name"}'))]
        )
        client = SimpleNamespace(
            chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(return_value=response)))
        )

        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", None)),
        ):
            assert await parser._identify_column_via_llm(df) == "Name"

        response = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content='{"target_column":"Missing"}'))]
        )
        client.chat.completions.create = AsyncMock(return_value=response)
        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", None)),
        ):
            assert await parser._identify_column_via_llm(df) is None

        with patch("matcher.ingestion.parser.llm_available", return_value=False):
            assert await parser._identify_column_via_llm(df) is None

        client.chat.completions.create = AsyncMock(side_effect=RuntimeError("llm down"))
        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch("matcher.ingestion.parser.make_llm_client", return_value=(client, "model", None)),
        ):
            assert await parser._identify_column_via_llm(df) is None

    def test_analyze_heuristic_splits_supplier_and_catalog_tables(self):
        df = pd.DataFrame(
            [
                ["Поставщик: ООО Ромашка", None, None, "Каталог", None],
                ["Номенклатура", "Ед. изм", None, "Наименование", "Цена"],
                ["Насос поставщика", "шт", None, "Насос каталог", "100"],
                ["Кабель поставщика", "м", None, "Кабель каталог", "200"],
            ]
        )

        result = parser._analyze_heuristic(df)

        assert result.tables_detected == 2
        assert result.supplier_name == "ООО Ромашка"
        assert [item["raw_text"] for item in result.supplier_items] == [
            "Насос поставщика",
            "Кабель поставщика",
        ]
        assert [item["raw_text"] for item in result.catalog_items] == [
            "Насос каталог",
            "Кабель каталог",
        ]

    def test_analyze_heuristic_uses_longest_text_when_no_name_column(self):
        df = pd.DataFrame(
            [
                ["x", "y"],
                ["A1", "Long descriptive product"],
                ["A2", "Other product"],
            ]
        )

        result = parser._analyze_heuristic(df)

        assert result.tables_detected == 1
        assert parser._analyze_heuristic(pd.DataFrame()).tables_detected == 0

        empty_text_df = pd.DataFrame([["x"], [None], [None]])
        result = parser._analyze_heuristic(empty_text_df)
        assert result.tables_detected == 1

        article_df = pd.DataFrame([["Артикул", "h"], ["A-1", "Medium"], ["B-2", "Longer"]])
        result = parser._analyze_heuristic(article_df)
        assert result.supplier_items[0]["raw_text"] == "Medium"

        with patch("matcher.ingestion.parser._split_into_column_groups", return_value=[(2, 1)]):
            assert parser._analyze_heuristic(pd.DataFrame([["x"]])).tables_detected == 1


class TestParserPublicApi:
    @pytest.mark.asyncio
    async def test_analyze_file_structure_uses_llm_tables_and_falls_back(self):
        contents = _xlsx_bytes(
            [
                ["Name", "Unit", None, "Catalog Name", "Price"],
                ["Pump", "pcs", None, "Catalog pump", "100"],
                ["Cable", "m", None, "Catalog cable", "200"],
            ]
        )
        llm_result = {
            "supplier_name": "Acme",
            "tables": [
                {
                    "role": "supplier_input",
                    "col_start": "A",
                    "col_end": "B",
                    "header_row": 1,
                    "name_col": "A",
                    "unit_col": "B",
                    "price_col": None,
                    "label": "supplier",
                },
                {
                    "role": "catalog_reference",
                    "col_start": "D",
                    "col_end": "E",
                    "header_row": 1,
                    "name_col": "D",
                    "unit_col": None,
                    "price_col": "E",
                    "label": "catalog",
                },
                {
                    "role": "metadata",
                    "col_start": "A",
                    "col_end": "B",
                    "header_row": 1,
                    "name_col": "A",
                },
                {"bad": "table"},
            ],
        }

        with patch("matcher.ingestion.parser._analyze_structure_via_llm", return_value=llm_result):
            result = await parser.analyze_file_structure(contents)

        assert result.supplier_name == "Acme"
        assert result.tables_detected == 4
        assert [item["raw_text"] for item in result.supplier_items] == ["Pump", "Cable"]
        assert [item["raw_text"] for item in result.catalog_items] == [
            "Catalog pump",
            "Catalog cable",
        ]

        with (
            patch(
                "matcher.ingestion.parser._get_best_sheet",
                return_value=(pd.DataFrame(), "Empty"),
            ),
            patch("matcher.ingestion.parser._analyze_structure_via_llm", new_callable=AsyncMock),
        ):
            assert await parser.analyze_file_structure(contents) == parser.FileAnalysisResult()

        with (
            patch(
                "matcher.ingestion.parser._get_best_sheet",
                return_value=(pd.DataFrame([["x"]]), "S"),
            ),
            patch("matcher.ingestion.parser._analyze_structure_via_llm", return_value=None),
            patch(
                "matcher.ingestion.parser._analyze_heuristic",
                return_value=parser.FileAnalysisResult(),
            ),
            patch("matcher.ingestion.parser.llm_available", return_value=False),
            patch(
                "matcher.ingestion.parser.parse_excel_upload",
                new_callable=AsyncMock,
                return_value=[{"raw_text": "Fallback", "line_id": "1", "original_row": {}}],
            ),
        ):
            result = await parser.analyze_file_structure(contents)

        assert result.tables_detected == 1
        assert result.supplier_items[0]["raw_text"] == "Fallback"

        heuristic_result = parser.FileAnalysisResult(
            supplier_items=[{"raw_text": "Heuristic", "line_id": "1", "original_row": {}}],
            tables_detected=1,
        )
        with (
            patch(
                "matcher.ingestion.parser._get_best_sheet",
                return_value=(pd.DataFrame([["Name"], ["Heuristic"]]), "S"),
            ),
            patch("matcher.ingestion.parser._analyze_structure_via_llm", return_value=None),
            patch("matcher.ingestion.parser._analyze_heuristic", return_value=heuristic_result),
        ):
            result = await parser.analyze_file_structure(contents)

        assert result.supplier_items[0]["raw_text"] == "Heuristic"

    @pytest.mark.asyncio
    async def test_parse_excel_upload_csv_heuristic_ai_and_empty_paths(self):
        csv_contents = "Номенклатура клиента;Цена\nPump;100\nPump;100\nИтого;200\nC;1\n".encode()

        items = await parser.parse_excel_upload(csv_contents)

        assert [item["raw_text"] for item in items] == ["Pump"]

        with patch("matcher.ingestion.parser.llm_available", return_value=False):
            with pytest.raises(ValueError, match="API ключ"):
                await parser.parse_excel_upload(csv_contents, use_ai=True)

        with (
            patch("matcher.ingestion.parser.llm_available", return_value=True),
            patch(
                "matcher.ingestion.parser._identify_column_via_llm", new_callable=AsyncMock
            ) as identify,
        ):
            identify.return_value = "Цена"
            items = await parser.parse_excel_upload(csv_contents, use_ai=True)

        assert items[0]["raw_text"] == "100"

        with patch("matcher.ingestion.parser._read_csv_raw", return_value=pd.DataFrame()):
            assert await parser.parse_excel_upload(b"") == []

    @pytest.mark.asyncio
    async def test_parse_excel_upload_excel_read_fallback_and_text_density(self):
        contents = _xlsx_bytes(
            [["Code", "Description"], ["1", "Very long product name"], ["2", "Other name"]]
        )

        with patch("matcher.ingestion.parser.pd.read_excel", side_effect=RuntimeError("bad xlsx")):
            items = await parser.parse_excel_upload(contents)

        assert [item["raw_text"] for item in items] == ["Very long product name", "Other name"]

        csv_contents = b"1,Long product text\n2,Other long product\n"
        items = await parser.parse_excel_upload(csv_contents)
        assert items[0]["raw_text"] == "Other long product"

        sparse_contents = b"1,\n2,\n"
        items = await parser.parse_excel_upload(sparse_contents)
        assert items == []
